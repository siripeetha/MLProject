# -*- coding: utf-8 -*-
"""
Resume Parser Service
Extracts text and contact metadata from PDF, DOCX, and TXT files.
"""

import os
import re
from pypdf import PdfReader
import docx


def extract_text_from_pdf(file_path):
    """Extract full text from a PDF document."""
    text_parts = []
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
    except Exception as e:
        print(f"Error reading PDF {file_path}: {e}")
    return "\n".join(text_parts).strip()


def extract_text_from_docx(file_path):
    """Extract full text from a Word DOCX document."""
    text_parts = []
    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            if p.text.strip():
                text_parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text.strip())
    except Exception as e:
        print(f"Error reading DOCX {file_path}: {e}")
    return "\n".join(text_parts).strip()


def extract_text_from_txt(file_path):
    """Extract text from a plain TXT file with fallback encodings."""
    for enc in ["utf-8", "latin-1", "cp1252"]:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read().strip()
        except UnicodeDecodeError:
            continue
    return ""


def parse_resume_file(file_path):
    """Extract text from PDF, DOCX, or TXT file."""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_path)
    elif ext == ".txt":
        return extract_text_from_txt(file_path)
    else:
        # Fallback to text reading
        return extract_text_from_txt(file_path)


def extract_contact_info(text, filename=""):
    """Heuristic extraction of candidate name, email, phone from resume text."""
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    
    # Email extraction
    email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', text)
    email = email_match.group(0) if email_match else "candidate@example.com"
    
    # Phone extraction
    phone_match = re.search(r'(\+?\d{1,3}[-.\s]?)?\(?\d{3,5}\)?[-.\s]?\d{3,5}[-.\s]?\d{3,5}', text)
    phone = phone_match.group(0).strip() if phone_match else "+91-98765-43210"
    
    # Candidate name extraction
    # Often the first non-contact line is the name
    name = "Candidate"
    for line in lines[:5]:
        if "@" in line or any(c.isdigit() for c in line) or len(line) > 50 or "resume" in line.lower() or "curriculum" in line.lower():
            continue
        words = line.split()
        if 1 <= len(words) <= 4:
            name = line.title()
            break
            
    if name == "Candidate" and filename:
        base = os.path.splitext(filename)[0]
        base_clean = re.sub(r'[^a-zA-Z\s]', ' ', base).strip()
        if base_clean:
            name = base_clean.title()
            
    return {
        "candidate_name": name,
        "email": email,
        "phone": phone
    }
